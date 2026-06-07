#!/usr/bin/env python3
"""Drive-style local upload/search POC backed by optional PaddleOCR."""

from __future__ import annotations

import cgi
import json
import mimetypes
import os
import queue
import shutil
import sqlite3
import threading
import time
import traceback
import uuid
from datetime import datetime, timezone
from http import HTTPStatus
from http.server import ThreadingHTTPServer, BaseHTTPRequestHandler
from pathlib import Path
from urllib.parse import parse_qs, unquote, urlparse

ROOT = Path(__file__).resolve().parent.parent
UPLOADS_DIR = ROOT / "uploads"
DATA_DIR = ROOT / "data"
DB_PATH = DATA_DIR / "paddle_drive.sqlite3"
STATIC_DIR = ROOT / "app" / "static"
PADDLEX_CACHE_DIR = DATA_DIR / "paddlex-cache"
os.environ.setdefault("PADDLE_PDX_CACHE_HOME", str(PADDLEX_CACHE_DIR))
os.environ.setdefault("XDG_CACHE_HOME", str(DATA_DIR / "cache"))
MAX_UPLOAD_BYTES = 50 * 1024 * 1024
SUPPORTED_IMAGE_MIMES = {
    "image/png",
    "image/jpeg",
    "image/webp",
    "image/bmp",
    "image/tiff",
}
SUPPORTED_IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}
SUPPORTED_WORD_MIMES = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}
SUPPORTED_EXCEL_MIMES = {
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
SUPPORTED_WORD_EXTS = {".docx"}
SUPPORTED_EXCEL_EXTS = {".xlsx"}

ocr_jobs: "queue.Queue[str]" = queue.Queue()
ocr_lock = threading.Lock()
ocr_engine = None
ocr_init_error = None


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def connect_db() -> sqlite3.Connection:
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_storage() -> None:
    UPLOADS_DIR.mkdir(exist_ok=True)
    DATA_DIR.mkdir(exist_ok=True)
    with connect_db() as conn:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS files (
                id TEXT PRIMARY KEY,
                original_name TEXT NOT NULL,
                stored_name TEXT NOT NULL,
                mime_type TEXT NOT NULL,
                size INTEGER NOT NULL,
                created_at TEXT NOT NULL,
                ocr_status TEXT NOT NULL,
                ocr_error TEXT
            )
            """
        )
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS ocr_text (
                file_id TEXT PRIMARY KEY,
                extracted_text TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                FOREIGN KEY(file_id) REFERENCES files(id) ON DELETE CASCADE
            )
            """
        )
        conn.commit()


def row_to_file(row: sqlite3.Row, extracted_text: str | None = None, snippet: str | None = None) -> dict:
    return {
        "id": row["id"],
        "original_name": row["original_name"],
        "mime_type": row["mime_type"],
        "size": row["size"],
        "created_at": row["created_at"],
        "ocr_status": row["ocr_status"],
        "ocr_error": row["ocr_error"],
        "download_url": f"/api/files/{row['id']}/download",
        "extracted_text": extracted_text or "",
        "snippet": snippet or "",
    }


def fetch_file(file_id: str) -> dict | None:
    with connect_db() as conn:
        row = conn.execute(
            """
            SELECT f.*, COALESCE(o.extracted_text, '') AS extracted_text
            FROM files f
            LEFT JOIN ocr_text o ON o.file_id = f.id
            WHERE f.id = ?
            """,
            (file_id,),
        ).fetchone()
    if not row:
        return None
    return row_to_file(row, row["extracted_text"])


def list_files(filter_name: str = "all") -> list[dict]:
    where = ""
    params: tuple = ()
    if filter_name == "indexed":
        where = "WHERE f.ocr_status = ?"
        params = ("indexed",)
    with connect_db() as conn:
        rows = conn.execute(
            f"""
            SELECT f.*, COALESCE(o.extracted_text, '') AS extracted_text
            FROM files f
            LEFT JOIN ocr_text o ON o.file_id = f.id
            {where}
            ORDER BY f.created_at DESC
            """,
            params,
        ).fetchall()
    return [row_to_file(row, row["extracted_text"]) for row in rows]


def save_uploaded_file(field) -> dict:
    original_name = Path(field.filename or "upload.bin").name
    mime_type = field.type or mimetypes.guess_type(original_name)[0] or "application/octet-stream"
    file_id = uuid.uuid4().hex
    suffix = Path(original_name).suffix.lower()
    stored_name = f"{file_id}{suffix or '.bin'}"
    target = UPLOADS_DIR / stored_name
    size = 0

    with target.open("wb") as out:
        while True:
            chunk = field.file.read(1024 * 1024)
            if not chunk:
                break
            size += len(chunk)
            if size > MAX_UPLOAD_BYTES:
                out.close()
                target.unlink(missing_ok=True)
                raise ValueError("Files must be 50 MB or smaller for this local POC.")
            out.write(chunk)

    status = "queued" if is_indexable_file(original_name, mime_type) else "unsupported"
    error = None if status == "queued" else "Search indexing currently supports images, DOCX, and XLSX files."
    with connect_db() as conn:
        conn.execute(
            """
            INSERT INTO files (id, original_name, stored_name, mime_type, size, created_at, ocr_status, ocr_error)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (file_id, original_name, stored_name, mime_type, size, now_iso(), status, error),
        )
        conn.commit()
    if status == "queued":
        ocr_jobs.put(file_id)
    return fetch_file(file_id) or {}


def is_supported_image(filename: str, mime_type: str) -> bool:
    return mime_type in SUPPORTED_IMAGE_MIMES or Path(filename).suffix.lower() in SUPPORTED_IMAGE_EXTS


def is_supported_word(filename: str, mime_type: str) -> bool:
    return mime_type in SUPPORTED_WORD_MIMES or Path(filename).suffix.lower() in SUPPORTED_WORD_EXTS


def is_supported_excel(filename: str, mime_type: str) -> bool:
    return mime_type in SUPPORTED_EXCEL_MIMES or Path(filename).suffix.lower() in SUPPORTED_EXCEL_EXTS


def is_indexable_file(filename: str, mime_type: str) -> bool:
    return (
        is_supported_image(filename, mime_type)
        or is_supported_word(filename, mime_type)
        or is_supported_excel(filename, mime_type)
    )


def set_ocr_status(file_id: str, status: str, error: str | None = None) -> None:
    with connect_db() as conn:
        conn.execute(
            "UPDATE files SET ocr_status = ?, ocr_error = ? WHERE id = ?",
            (status, error, file_id),
        )
        conn.commit()


def get_ocr_engine():
    global ocr_engine, ocr_init_error
    with ocr_lock:
        if ocr_engine is not None:
            return ocr_engine
        if ocr_init_error:
            raise RuntimeError(ocr_init_error)
        try:
            from paddleocr import PaddleOCR  # type: ignore

            for kwargs in (
                {"use_textline_orientation": True, "lang": "en"},
                {"use_angle_cls": True, "lang": "en"},
                {"lang": "en"},
            ):
                try:
                    ocr_engine = PaddleOCR(**kwargs)
                    break
                except (TypeError, ValueError):
                    continue
            if ocr_engine is None:
                raise RuntimeError("No compatible PaddleOCR constructor signature was accepted.")
            return ocr_engine
        except Exception as exc:  # pragma: no cover - depends on local install/models
            ocr_init_error = (
                "PaddleOCR could not be initialized. Install dependencies with "
                "`python3 -m pip install -r requirements.txt`, then restart. "
                f"Details: {exc}"
            )
            raise RuntimeError(ocr_init_error) from exc


def flatten_ocr_result(result) -> str:
    texts: list[str] = []

    def walk(node) -> None:
        if node is None:
            return
        if isinstance(node, dict):
            for key in ("rec_texts", "texts"):
                value = node.get(key)
                if isinstance(value, list):
                    texts.extend(str(v) for v in value if str(v).strip())
            for value in node.values():
                walk(value)
            return
        if isinstance(node, (list, tuple)):
            if len(node) >= 2 and isinstance(node[1], (list, tuple)) and node[1]:
                maybe_text = node[1][0]
                if isinstance(maybe_text, str) and maybe_text.strip():
                    texts.append(maybe_text)
            for item in node:
                walk(item)

    walk(result)
    seen = set()
    unique = []
    for text in texts:
        clean = " ".join(text.split())
        if clean and clean not in seen:
            seen.add(clean)
            unique.append(clean)
    return "\n".join(unique)


def extract_docx_text(file_path: Path) -> str:
    from docx import Document  # type: ignore

    doc = Document(str(file_path))
    parts: list[str] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [" ".join(cell.text.split()) for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_xlsx_text(file_path: Path) -> str:
    from openpyxl import load_workbook  # type: ignore

    workbook = load_workbook(filename=str(file_path), read_only=True, data_only=True)
    parts: list[str] = []
    try:
        for sheet in workbook.worksheets:
            parts.append(f"Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = [str(value).strip() for value in row if value is not None and str(value).strip()]
                if values:
                    parts.append(" | ".join(values))
    finally:
        workbook.close()
    return "\n".join(parts)


def store_extracted_text(file_id: str, extracted: str) -> None:
    with connect_db() as conn:
        conn.execute(
            """
            INSERT INTO ocr_text (file_id, extracted_text, updated_at)
            VALUES (?, ?, ?)
            ON CONFLICT(file_id) DO UPDATE SET
                extracted_text = excluded.extracted_text,
                updated_at = excluded.updated_at
            """,
            (file_id, extracted, now_iso()),
        )
        conn.execute(
            "UPDATE files SET ocr_status = ?, ocr_error = NULL WHERE id = ?",
            ("indexed", file_id),
        )
        conn.commit()


def run_ocr(file_id: str) -> None:
    with connect_db() as conn:
        row = conn.execute("SELECT * FROM files WHERE id = ?", (file_id,)).fetchone()
    if not row:
        return
    if not is_indexable_file(row["original_name"], row["mime_type"]):
        set_ocr_status(file_id, "unsupported", "Search indexing currently supports images, DOCX, and XLSX files.")
        return

    set_ocr_status(file_id, "processing")
    file_path = UPLOADS_DIR / row["stored_name"]
    try:
        if is_supported_word(row["original_name"], row["mime_type"]):
            extracted = extract_docx_text(file_path)
        elif is_supported_excel(row["original_name"], row["mime_type"]):
            extracted = extract_xlsx_text(file_path)
        else:
            engine = get_ocr_engine()
            if hasattr(engine, "predict"):
                result = engine.predict(str(file_path))
            else:
                result = engine.ocr(str(file_path))
            extracted = flatten_ocr_result(result)
        store_extracted_text(file_id, extracted.strip())
    except Exception as exc:  # pragma: no cover - depends on local install/models
        set_ocr_status(file_id, "error", f"{exc}")
        traceback.print_exc()


def ocr_worker() -> None:
    while True:
        file_id = ocr_jobs.get()
        try:
            run_ocr(file_id)
        finally:
            ocr_jobs.task_done()


def make_snippet(text: str, query: str) -> str:
    if not text:
        return ""
    lower = text.lower()
    q = query.lower()
    index = lower.find(q)
    if index < 0:
        return text[:180]
    start = max(0, index - 70)
    end = min(len(text), index + len(query) + 110)
    prefix = "..." if start > 0 else ""
    suffix = "..." if end < len(text) else ""
    return f"{prefix}{text[start:end]}{suffix}"


def search_files(query: str) -> list[dict]:
    q = query.strip()
    if not q:
        return list_files()
    like = f"%{q}%"
    with connect_db() as conn:
        rows = conn.execute(
            """
            SELECT f.*, COALESCE(o.extracted_text, '') AS extracted_text,
                CASE
                    WHEN LOWER(f.original_name) = LOWER(?) THEN 0
                    WHEN LOWER(f.original_name) LIKE LOWER(?) THEN 1
                    WHEN LOWER(COALESCE(o.extracted_text, '')) LIKE LOWER(?) THEN 2
                    ELSE 3
                END AS rank
            FROM files f
            LEFT JOIN ocr_text o ON o.file_id = f.id
            WHERE LOWER(f.original_name) LIKE LOWER(?)
               OR LOWER(COALESCE(o.extracted_text, '')) LIKE LOWER(?)
            ORDER BY rank ASC, f.created_at DESC
            """,
            (q, like, like, like, like),
        ).fetchall()
    return [row_to_file(row, row["extracted_text"], make_snippet(row["extracted_text"], q)) for row in rows]


class Handler(BaseHTTPRequestHandler):
    server_version = "PaddleDrivePOC/0.1"

    def do_GET(self) -> None:
        parsed = urlparse(self.path)
        path = parsed.path
        if path == "/":
            self.serve_static("index.html", "text/html; charset=utf-8")
        elif path == "/app.js":
            self.serve_static("app.js", "application/javascript; charset=utf-8")
        elif path == "/styles.css":
            self.serve_static("styles.css", "text/css; charset=utf-8")
        elif path == "/api/files":
            filter_name = parse_qs(parsed.query).get("filter", ["all"])[0]
            self.send_json({"files": list_files(filter_name)})
        elif path == "/api/search":
            q = parse_qs(parsed.query).get("q", [""])[0]
            self.send_json({"files": search_files(q), "query": q})
        elif path.startswith("/api/files/") and path.endswith("/download"):
            self.download_file(path)
        else:
            self.send_error(HTTPStatus.NOT_FOUND)

    def do_POST(self) -> None:
        parsed = urlparse(self.path)
        if parsed.path == "/api/files":
            self.handle_upload()
        elif parsed.path.startswith("/api/files/") and parsed.path.endswith("/ocr"):
            file_id = parsed.path.split("/")[3]
            file_info = fetch_file(file_id)
            if not file_info:
                self.send_error(HTTPStatus.NOT_FOUND)
                return
            set_ocr_status(file_id, "queued")
            ocr_jobs.put(file_id)
            self.send_json({"file": fetch_file(file_id)})
        else:
            self.send_error(HTTPStatus.NOT_FOUND)

    def do_DELETE(self) -> None:
        parsed = urlparse(self.path)
        if parsed.path.startswith("/api/files/"):
            file_id = parsed.path.split("/")[3]
            with connect_db() as conn:
                row = conn.execute("SELECT * FROM files WHERE id = ?", (file_id,)).fetchone()
                if not row:
                    self.send_error(HTTPStatus.NOT_FOUND)
                    return
                (UPLOADS_DIR / row["stored_name"]).unlink(missing_ok=True)
                conn.execute("DELETE FROM ocr_text WHERE file_id = ?", (file_id,))
                conn.execute("DELETE FROM files WHERE id = ?", (file_id,))
                conn.commit()
            self.send_json({"ok": True})
        else:
            self.send_error(HTTPStatus.NOT_FOUND)

    def handle_upload(self) -> None:
        ctype, pdict = cgi.parse_header(self.headers.get("content-type"))
        if ctype != "multipart/form-data":
            self.send_error(HTTPStatus.BAD_REQUEST, "Expected multipart/form-data.")
            return
        pdict["boundary"] = bytes(pdict["boundary"], "utf-8")
        form = cgi.FieldStorage(fp=self.rfile, headers=self.headers, environ={"REQUEST_METHOD": "POST"})
        fields = form["files"] if "files" in form else []
        if not isinstance(fields, list):
            fields = [fields]
        saved = []
        try:
            for field in fields:
                if getattr(field, "filename", None):
                    saved.append(save_uploaded_file(field))
        except ValueError as exc:
            self.send_error(HTTPStatus.BAD_REQUEST, str(exc))
            return
        self.send_json({"files": saved}, HTTPStatus.CREATED)

    def download_file(self, path: str) -> None:
        file_id = path.split("/")[3]
        with connect_db() as conn:
            row = conn.execute("SELECT * FROM files WHERE id = ?", (file_id,)).fetchone()
        if not row:
            self.send_error(HTTPStatus.NOT_FOUND)
            return
        file_path = UPLOADS_DIR / row["stored_name"]
        if not file_path.exists():
            self.send_error(HTTPStatus.NOT_FOUND)
            return
        self.send_response(HTTPStatus.OK)
        self.send_header("Content-Type", row["mime_type"])
        self.send_header("Content-Length", str(file_path.stat().st_size))
        self.send_header("Content-Disposition", f"attachment; filename*=UTF-8''{row['original_name']}")
        self.end_headers()
        with file_path.open("rb") as fh:
            shutil.copyfileobj(fh, self.wfile)

    def serve_static(self, filename: str, content_type: str) -> None:
        path = STATIC_DIR / filename
        if not path.exists():
            self.send_error(HTTPStatus.NOT_FOUND)
            return
        data = path.read_bytes()
        self.send_response(HTTPStatus.OK)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def send_json(self, payload: dict, status: HTTPStatus = HTTPStatus.OK) -> None:
        data = json.dumps(payload).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def log_message(self, fmt: str, *args) -> None:
        print(f"[{time.strftime('%H:%M:%S')}] {fmt % args}")


def main() -> None:
    init_storage()
    threading.Thread(target=ocr_worker, daemon=True).start()
    host = os.environ.get("HOST", "127.0.0.1")
    port = int(os.environ.get("PORT", "8765"))
    server = ThreadingHTTPServer((host, port), Handler)
    print(f"Paddle Drive POC running at http://{host}:{port}")
    print("Upload images to trigger OCR. Press Ctrl+C to stop.")
    server.serve_forever()


if __name__ == "__main__":
    main()
